import tkinter as tk
from tkinter import filedialog, messagebox
import os
import platform
import subprocess
from pptx import Presentation
import copy
import re
import json

# 전역 변수: 추가된 노래 데이터를 저장할 리스트
songs_data = []

def save_song_list():
    """현재 추가된 노래 목록을 파일(.json)로 저장합니다."""
    if not songs_data:
        messagebox.showwarning("경고", "저장할 노래가 목록에 없습니다.")
        return

    save_path = filedialog.asksaveasfilename(
        title="노래 목록 저장",
        defaultextension=".json",
        filetypes=[("JSON 파일", "*.json"), ("모든 파일", "*.*")]
    )

    if save_path:
        try:
            # json 형식으로 안전하게 저장 (한글 깨짐 방지를 위해 ensure_ascii=False)
            with open(save_path, 'w', encoding='utf-8') as f:
                json.dump(songs_data, f, ensure_ascii=False, indent=4)
            lbl_status.config(text="노래 목록이 저장되었습니다.", fg="blue")
            messagebox.showinfo("성공", "목록이 성공적으로 저장되었습니다.")
        except Exception as e:
            messagebox.showerror("오류", f"저장 중 오류가 발생했습니다.\n{str(e)}")

def load_song_list():
    """검색, 장바구니, DB 원본 수정 및 삭제 기능이 모두 탑재된 만능 불러오기 창"""
    open_path = filedialog.askopenfilename(
        title="노래 데이터베이스 불러오기",
        filetypes=[("JSON 파일", "*.json"), ("모든 파일", "*.*")]
    )

    if not open_path:
        return

    try:
        with open(open_path, 'r', encoding='utf-8') as f:
            loaded_data = json.load(f)
    except Exception as e:
        messagebox.showerror("오류", f"파일을 불러오는 중 오류가 발생했습니다.\n{str(e)}")
        return

    # 1. 팝업창 생성
    popup = tk.Toplevel(root)
    popup.title("노래 검색 및 DB 관리 / 장바구니")
    popup.geometry("750x600") 
    popup.grab_set() 

    frame_main = tk.Frame(popup)
    frame_main.pack(fill="both", expand=True, padx=15, pady=15)

    # ==========================================
    # [왼쪽 영역] 검색 및 DB 목록
    # ==========================================
    frame_left = tk.Frame(frame_main)
    frame_left.pack(side="left", fill="both", expand=True)

    tk.Label(frame_left, text="🔍 제목 검색 및 DB 관리", font=TITLE_FONT).pack(anchor="w")
    search_var = tk.StringVar()
    entry_search = tk.Entry(frame_left, textvariable=search_var, font=NORMAL_FONT)
    entry_search.pack(fill="x", pady=(0, 5))

    frame_db_list = tk.Frame(frame_left)
    frame_db_list.pack(fill="both", expand=True)
    
    scroll_db = tk.Scrollbar(frame_db_list)
    scroll_db.pack(side="right", fill="y")
    
    listbox_db = tk.Listbox(frame_db_list, yscrollcommand=scroll_db.set, font=NORMAL_FONT, selectmode=tk.EXTENDED)
    listbox_db.pack(fill="both", expand=True)
    scroll_db.config(command=listbox_db.yview)

    displayed_songs = []

    def update_db_list(*args):
        search_keyword = search_var.get().lower().replace(" ", "")
        listbox_db.delete(0, tk.END)
        displayed_songs.clear()
        
        for song in loaded_data:
            target_title = song['title'].lower().replace(" ", "")
            if not search_keyword or search_keyword in target_title:
                listbox_db.insert(tk.END, song['title'])
                displayed_songs.append(song)

    # 🎯 [새로 추가된 기능] DB 원본 수정하기
    def edit_db_song():
        selected_indices = listbox_db.curselection()
        if len(selected_indices) != 1:
            messagebox.showwarning("알림", "수정할 곡을 하나만 선택해주세요.", parent=popup)
            return
        
        idx = selected_indices[0]
        song_to_edit = displayed_songs[idx]
        original_idx = loaded_data.index(song_to_edit) # 원본 데이터에서의 위치 찾기
        
        # 수정용 미니 팝업창 띄우기
        edit_popup = tk.Toplevel(popup)
        edit_popup.title("DB 곡 수정")
        edit_popup.geometry("400x500")
        edit_popup.grab_set()
        
        tk.Label(edit_popup, text="수정할 제목:", font=NORMAL_FONT).pack(anchor="w", padx=10, pady=(10, 0))
        entry_title = tk.Entry(edit_popup, font=NORMAL_FONT)
        entry_title.pack(fill="x", padx=10, pady=5)
        entry_title.insert(0, song_to_edit['title'])
        
        tk.Label(edit_popup, text="수정할 가사:", font=NORMAL_FONT).pack(anchor="w", padx=10)
        text_lyrics = tk.Text(edit_popup, font=NORMAL_FONT, height=15)
        text_lyrics.pack(fill="both", expand=True, padx=10, pady=5)
        text_lyrics.insert("1.0", song_to_edit['lyrics'])
        
        def save_edit():
            new_title = entry_title.get().strip()
            new_lyrics = text_lyrics.get("1.0", tk.END).strip()
            
            if not new_title or not new_lyrics:
                messagebox.showwarning("알림", "제목과 가사를 모두 입력해주세요.", parent=edit_popup)
                return
            
            # 원본 데이터 업데이트
            loaded_data[original_idx]['title'] = new_title
            loaded_data[original_idx]['lyrics'] = new_lyrics
            
            # JSON 파일 덮어쓰기
            try:
                with open(open_path, 'w', encoding='utf-8') as f:
                    json.dump(loaded_data, f, ensure_ascii=False, indent=4)
                messagebox.showinfo("성공", "DB 내용이 수정되었습니다.", parent=edit_popup)
                edit_popup.destroy()
                update_db_list() # 목록 새로고침
            except Exception as e:
                messagebox.showerror("오류", f"저장 실패:\n{e}", parent=edit_popup)
                
        tk.Button(edit_popup, text="✔ 내용 DB에 저장", bg="#2196F3", fg="white", 
                  font=("맑은 고딕", 11, "bold"), height=2, command=save_edit).pack(fill="x", padx=10, pady=10)

    # 🎯 [새로 추가된 기능] DB 원본 삭제하기
    def delete_db_song():
        selected_indices = listbox_db.curselection()
        if not selected_indices:
            messagebox.showwarning("알림", "삭제할 곡을 선택해주세요.", parent=popup)
            return
            
        if messagebox.askyesno("확인", f"선택한 {len(selected_indices)}곡을 DB에서 영구 삭제하시겠습니까?", parent=popup):
            # 뒤에서부터 삭제해야 인덱스가 꼬이지 않음
            for idx in reversed(selected_indices):
                song_to_delete = displayed_songs[idx]
                loaded_data.remove(song_to_delete)
                
            try:
                with open(open_path, 'w', encoding='utf-8') as f:
                    json.dump(loaded_data, f, ensure_ascii=False, indent=4)
                update_db_list() # 목록 새로고침
            except Exception as e:
                messagebox.showerror("오류", f"삭제 실패:\n{e}", parent=popup)

    # 🎯 왼쪽 영역 하단: 수정/삭제 버튼 배치
    frame_db_btns = tk.Frame(frame_left)
    frame_db_btns.pack(fill="x", pady=(5, 0))
    
    btn_edit_db = tk.Button(frame_db_btns, text="✏️ DB 곡 수정", font=("맑은 고딕", 9), command=edit_db_song)
    btn_edit_db.pack(side="left", fill="x", expand=True, padx=(0, 2))
    
    btn_delete_db = tk.Button(frame_db_btns, text="🗑️ DB 영구 삭제", font=("맑은 고딕", 9), command=delete_db_song)
    btn_delete_db.pack(side="right", fill="x", expand=True, padx=(2, 0))

    # ==========================================
    # [가운데 영역] 담기 / 빼기 버튼 (기존과 동일)
    # ==========================================
    frame_mid = tk.Frame(frame_main)
    frame_mid.pack(side="left", fill="y", padx=10)

    cart_data = []

    def add_to_cart():
        selected_indices = listbox_db.curselection()
        for idx in selected_indices:
            cart_data.append(displayed_songs[idx])
        refresh_cart_listbox()
        listbox_db.selection_clear(0, tk.END)
        search_var.set("")
        entry_search.focus()

    def remove_from_cart():
        selected_indices = listbox_cart.curselection()
        for idx in reversed(selected_indices):
            del cart_data[idx]
        refresh_cart_listbox()

    def refresh_cart_listbox():
        listbox_cart.delete(0, tk.END)
        for i, song in enumerate(cart_data):
            listbox_cart.insert(tk.END, f"{i+1}. {song['title']}")
        lbl_cart_status.config(text=f"현재 담긴 곡: {len(cart_data)}곡")

    tk.Label(frame_mid, text="").pack(pady=70)
    btn_add = tk.Button(frame_mid, text="담기 ▶", bg="#4CAF50", fg="white", font=NORMAL_FONT, command=add_to_cart)
    btn_add.pack(fill="x", pady=5)
    btn_remove = tk.Button(frame_mid, text="◀ 빼기", bg="#f44336", fg="white", font=NORMAL_FONT, command=remove_from_cart)
    btn_remove.pack(fill="x", pady=5)

    # ==========================================
    # [오른쪽 영역] 장바구니 (담은 목록) (기존과 동일)
    # ==========================================
    frame_right = tk.Frame(frame_main)
    frame_right.pack(side="right", fill="both", expand=True)

    tk.Label(frame_right, text="🛒 담은 목록 (장바구니)", font=TITLE_FONT).pack(anchor="w")
    lbl_cart_status = tk.Label(frame_right, text="현재 담긴 곡: 0곡", fg="blue", font=NORMAL_FONT)
    lbl_cart_status.pack(anchor="w", pady=(0, 4))

    frame_cart_list = tk.Frame(frame_right)
    frame_cart_list.pack(fill="both", expand=True)
    
    scroll_cart = tk.Scrollbar(frame_cart_list)
    scroll_cart.pack(side="right", fill="y")
    
    listbox_cart = tk.Listbox(frame_cart_list, yscrollcommand=scroll_cart.set, font=NORMAL_FONT, selectmode=tk.EXTENDED)
    listbox_cart.pack(fill="both", expand=True)
    scroll_cart.config(command=listbox_cart.yview)

    search_var.trace_add("write", update_db_list)
    update_db_list()

    # ==========================================
    # [하단 영역] 최종 적용 버튼
    # ==========================================
    def confirm_and_close():
        if not cart_data:
            popup.destroy()
            return
            
        for song in cart_data:
            songs_data.append(song)
            listbox_songs.insert(tk.END, f"{len(songs_data)}. {song['title']}")
            
        lbl_status.config(text=f"장바구니에서 {len(cart_data)}곡을 성공적으로 추가했습니다.", fg="green")
        popup.destroy()

    frame_bottom = tk.Frame(popup)
    frame_bottom.pack(fill="x", padx=15, pady=10)
    
    btn_confirm = tk.Button(frame_bottom, text="✔ 장바구니 곡들 최종 추가하기", bg="#2196F3", fg="white", 
                            font=("맑은 고딕", 12, "bold"), height=2, command=confirm_and_close)
    btn_confirm.pack(fill="x")

def append_to_song_list():
    """현재 화면에 있는 노래들을 기존에 저장된 파일(.json)에 누적해서 추가합니다."""
    if not songs_data:
        messagebox.showwarning("경고", "추가할 노래가 목록에 없습니다.")
        return

    # 1. 곡을 추가해 넣을 기존 파일 선택
    open_path = filedialog.askopenfilename(
        title="어느 파일에 곡을 추가할까요? (기존 DB 선택)",
        filetypes=[("JSON 파일", "*.json"), ("모든 파일", "*.*")]
    )

    if not open_path:
        return

    # 2. 기존 파일에 있던 노래 데이터 읽어오기
    try:
        with open(open_path, 'r', encoding='utf-8') as f:
            existing_data = json.load(f)
    except Exception as e:
        messagebox.showerror("오류", f"기존 파일을 읽는 중 오류가 발생했습니다.\n{str(e)}")
        return

    # 3. 기존 데이터 뒤에 현재 화면의 데이터(새 곡) 이어 붙이기
    existing_data.extend(songs_data)

    # 4. 곡이 합쳐진 데이터를 같은 파일에 다시 덮어쓰기(저장)
    try:
        with open(open_path, 'w', encoding='utf-8') as f:
            json.dump(existing_data, f, ensure_ascii=False, indent=4)
        lbl_status.config(text=f"기존 파일에 {len(songs_data)}곡이 추가되었습니다.", fg="blue")
        messagebox.showinfo("성공", f"성공적으로 누적 추가되었습니다!\n(해당 파일의 총 곡 수: {len(existing_data)}곡)")
    except Exception as e:
        messagebox.showerror("오류", f"파일을 저장하는 중 오류가 발생했습니다.\n{str(e)}")

def _copy_textbox_and_format(slide, ref_shape, new_text):
    """기존 텍스트 박스의 위치, 한글 폰트, 크기 등 모든 XML을 통째로 복제한 뒤 텍스트만 바꿉니다."""
    
    # 1. 텍스트 박스를 XML 수준에서 통째로 복제하여 새 슬라이드에 붙여넣기
    cloned_shape_element = copy.deepcopy(ref_shape.element)
    slide.shapes._spTree.append(cloned_shape_element)
    
    # 방금 복제되어 생성된 새 텍스트 박스 가져오기
    new_shape = slide.shapes[-1]
    
    # 2. 원본에서 서식(단락 정렬, 폰트/크기 등) 정보 백업해두기
    ref_pPr = None # 단락 서식 백업
    ref_rPr = None # 글자(폰트) 서식 백업
    
    if ref_shape.text_frame.paragraphs:
        ref_p = ref_shape.text_frame.paragraphs[0]
        if ref_p._p.pPr is not None:
            ref_pPr = copy.deepcopy(ref_p._p.pPr)
            
        for r in ref_p.runs:
            # 빈칸이 아닌 실제 글자가 있는 부분의 폰트(XML)를 백업
            if r.text.strip() and r._r.rPr is not None:
                ref_rPr = copy.deepcopy(r._r.rPr)
                break

    # 3. 텍스트 내용 덮어쓰기 (이때 파이썬이 임의로 서식을 초기화시켜 버림)
    new_shape.text_frame.text = new_text
    
    # 4. 초기화된 서식을 아까 백업해둔 진짜 원본 서식으로 강제 덮어쓰기
    for p in new_shape.text_frame.paragraphs:
        # 단락 서식 복구 (가운데 정렬, 줄 간격 등)
        if ref_pPr is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, copy.deepcopy(ref_pPr))
            
        # 글자 서식 복구 (수다공주 폰트, 글자 크기, 색상 등)
        for r in p.runs:
            if ref_rPr is not None:
                if r._r.rPr is not None:
                    r._r.remove(r._r.rPr)
                r._r.insert(0, copy.deepcopy(ref_rPr))

# ---------------------------------------------------------
# 1. 여러 곡 PPT 생성 로직 (제목/가사 자동 분리 적용)
# ---------------------------------------------------------
def generate_multiple_lyrics_ppt(template_path, songs_list, save_path, ref_slide_num, start_slide_num):
    try:
        prs = Presentation(template_path)
        
        if ref_slide_num < 1 or ref_slide_num > len(prs.slides):
            return False, f"템플릿에 {ref_slide_num}번 슬라이드가 없습니다."
        
        ref_slide = prs.slides[ref_slide_num - 1]
        
        # 1. 선택한 슬라이드에서 글자가 있는 '텍스트 박스'만 모두 찾기
        text_shapes = []
        for shape in ref_slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                text_shapes.append(shape)
                
        # 2. Y좌표(top)를 기준으로 오름차순 정렬 (화면 위에서부터 아래 순서로)
        text_shapes.sort(key=lambda s: s.top)
        
        if len(text_shapes) < 2:
            return False, "선택한 기준 슬라이드에 텍스트 박스가 최소 2개(위: 제목용, 아래: 가사용) 있어야 합니다."
            
        # 3. 위쪽에 있는 건 제목, 아래쪽에 있는 건 가사로 지정!
        title_ref = text_shapes[0]
        lyrics_ref = text_shapes[1]
        
        # 4. 그림(배경)을 그대로 가져오기 위해 기준 슬라이드의 '레이아웃'을 사용
        ref_layout = ref_slide.slide_layout

        # 가사가 들어갈 실제 시작 인덱스 (파이썬은 0부터 시작하므로 -1)
        insert_index = max(0, start_slide_num - 1)
        
        # 만약 입력한 위치가 기존 PPT 슬라이드 장수보다 크면 맨 끝에 붙이도록 방어 코드 추가
        if insert_index > len(prs.slides):
            insert_index = len(prs.slides)
        
        # 5. 노래 목록 돌면서 슬라이드 생성
        for song in songs_list:
            title = song["title"]
            lyrics = song["lyrics"]
            
            # 🎯 [수정된 코드] 오직 '빈 줄(엔터 2번)'만 슬라이드 분할 기준으로 삼습니다.
            lyric_chunks = re.split(r'\n\s*\n', lyrics.strip())
            
            for chunk in lyric_chunks:
                if not chunk.strip(): 
                    continue
                    
                # 1. 새 슬라이드 추가 
                new_slide = prs.slides.add_slide(ref_layout)
                
                # 방해되는 기본 틀 삭제
                for shape in new_slide.placeholders:
                    sp = shape.element
                    sp.getparent().remove(sp)
                
                # 2. 서식 복사 및 글자 넣기 
                # 🎯 [핵심] 프로그램이 임의로 손대지 않고, 사용자가 입력한 chunk를 그대로 넣습니다!
                _copy_textbox_and_format(new_slide, title_ref, title)
                _copy_textbox_and_format(new_slide, lyrics_ref, chunk.strip())
                
                # 3. 생성된 슬라이드를 원하는 위치로 이동
                slide_id_list = prs.slides._sldIdLst
                slide_to_move = slide_id_list[-1]
                slide_id_list.remove(slide_to_move)
                slide_id_list.insert(insert_index, slide_to_move)
                
                insert_index += 1
                
        prs.save(save_path)
        return True, save_path
        
    except Exception as e:
        return False, str(e)

# ---------------------------------------------------------
# 2. UI 동작 함수들
# ---------------------------------------------------------
def select_template():
    """템플릿 선택 창"""
    file_path = filedialog.askopenfilename(
        title="템플릿 PPT 선택",
        filetypes=[("PowerPoint Files", "*.pptx")]
    )
    if file_path:
        entry_template_path.config(state='normal')
        entry_template_path.delete(0, tk.END)
        entry_template_path.insert(0, file_path)
        entry_template_path.config(state='readonly')

def add_song():
    """입력한 노래를 목록에 추가"""
    title = entry_title.get().strip()
    lyrics = text_lyrics.get("1.0", tk.END).strip()

    if not title or not lyrics:
        messagebox.showwarning("입력 오류", "노래 제목과 가사를 모두 입력해주세요.")
        return

    # 데이터 저장 및 리스트박스에 표시
    songs_data.append({"title": title, "lyrics": lyrics})
    listbox_songs.insert(tk.END, f"{len(songs_data)}. {title}")

    # 다음 곡 입력을 위해 입력창 비우기
    entry_title.delete(0, tk.END)
    text_lyrics.delete("1.0", tk.END)
    
    lbl_status.config(text=f"'{title}' 곡이 목록에 추가되었습니다.", fg="blue")

def refresh_main_listbox():
    """메인 리스트박스의 번호와 목록을 깔끔하게 다시 매겨줍니다."""
    listbox_songs.delete(0, tk.END)
    for i, song in enumerate(songs_data):
        listbox_songs.insert(tk.END, f"{i+1}. {song['title']}")

def remove_song():
    """선택한 곡을 삭제하고 번호를 새로고침합니다."""
    selected_indices = listbox_songs.curselection()
    if not selected_indices:
        messagebox.showwarning("경고", "삭제할 곡을 선택하세요.")
        return
    
    # 인덱스가 꼬이지 않도록 뒤에서부터 삭제
    for idx in reversed(selected_indices):
        del songs_data[idx]
    
    refresh_main_listbox()
    lbl_status.config(text="선택한 곡이 삭제되었습니다.", fg="blue")

def move_song_up():
    """선택한 곡의 순서를 한 칸 위로 올립니다."""
    selected_indices = listbox_songs.curselection()
    if not selected_indices:
        return
        
    new_selection = []
    for idx in selected_indices:
        if idx > 0 and (idx - 1) not in new_selection:
            # 실제 데이터 위치 맞교환
            songs_data[idx - 1], songs_data[idx] = songs_data[idx], songs_data[idx - 1]
            new_selection.append(idx - 1)
        else:
            new_selection.append(idx)
            
    refresh_main_listbox()
    
    # 순서를 바꾼 뒤에도 선택(파란색 하이라이트) 상태 유지
    for idx in new_selection:
        listbox_songs.selection_set(idx)

def move_song_down():
    """선택한 곡의 순서를 한 칸 아래로 내립니다."""
    selected_indices = listbox_songs.curselection()
    if not selected_indices:
        return
        
    new_selection = []
    # 아래로 내릴 때는 뒤에서부터 계산해야 꼬이지 않음
    for idx in reversed(selected_indices):
        if idx < len(songs_data) - 1 and (idx + 1) not in new_selection:
            songs_data[idx + 1], songs_data[idx] = songs_data[idx], songs_data[idx + 1]
            new_selection.insert(0, idx + 1)
        else:
            new_selection.insert(0, idx)
            
    refresh_main_listbox()
    
    for idx in new_selection:
        listbox_songs.selection_set(idx)

def on_click_generate():
    """PPT 만들기 및 저장"""
    template = entry_template_path.get()

    if not template:
        messagebox.showwarning("입력 오류", "템플릿 PPT 파일을 선택해주세요.")
        return
    if len(songs_data) == 0:
        messagebox.showwarning("입력 오류", "목록에 추가된 노래가 없습니다.")
        return


    # 첫 번째 곡 제목을 활용해 기본 저장 파일명 생성
    default_filename = f"{songs_data[0]['title']} 외 {len(songs_data)-1}곡.pptx"
    
    save_path = filedialog.asksaveasfilename(
        title="완성된 PPT 저장 위치 선택",
        initialfile=default_filename,
        defaultextension=".pptx",
        filetypes=[("PowerPoint Files", "*.pptx")]
    )

    if not save_path:
        messagebox.showwarning("저장 오류", "저장할 파일 경로를 선택해주세요.")
        return

    lbl_status.config(text="PPT를 생성하는 중입니다...", fg="blue")
    root.update()

    # 입력된 슬라이드 번호 가져오기
    try:
        ref_slide_num = int(spin_slide.get())
        start_slide_num = int(spin_insert.get()) # 새롭게 추가된 시작 위치 번호
    except ValueError:
        messagebox.showwarning("입력 오류", "슬라이드 번호는 숫자로 입력해주세요.")
        return
    
    # 다중 곡 생성 함수 호출
    success, result = generate_multiple_lyrics_ppt(template, songs_data, save_path, ref_slide_num, start_slide_num)

    if success:
        lbl_status.config(text="생성 완료!", fg="green")
        msg_result = messagebox.askyesno(
            "성공", 
            f"총 {len(songs_data)}곡이 성공적으로 저장되었습니다!\n\n저장된 폴더를 여시겠습니까?"
        )
        if msg_result:
            folder_path = os.path.dirname(result)
            try:
                if platform.system() == "Windows":
                    os.startfile(folder_path)
                elif platform.system() == "Darwin":
                    subprocess.call(["open", folder_path])
            except Exception:
                pass
    else:
        messagebox.showerror("오류", f"PPT 생성 중 오류가 발생했습니다.\n{result}")
        lbl_status.config(text="오류 발생", fg="red")


# ---------------------------------------------------------
# 3. UI 화면 구성
# ---------------------------------------------------------
root = tk.Tk()
root.title("LyricSlide - 다중 곡 PPT 가사 자동 생성기")
root.geometry("600x750") # 창 크기를 조금 더 키움
root.resizable(False, False)

TITLE_FONT = ("맑은 고딕", 10, "bold")
NORMAL_FONT = ("맑은 고딕", 10)

# (1) 템플릿 선택
frame_template = tk.Frame(root)
frame_template.pack(fill="x", padx=20, pady=(15, 5))
tk.Label(frame_template, text="1. 기존 PPT (템플릿) 선택", font=TITLE_FONT).pack(anchor="w")
entry_template_path = tk.Entry(frame_template, width=55, state='readonly', font=NORMAL_FONT)
entry_template_path.pack(side="left", pady=5)
btn_file = tk.Button(frame_template, text="파일 찾기", command=select_template, font=NORMAL_FONT)
btn_file.pack(side="right", padx=(5, 0))

# 서식 복사할 슬라이드 번호 입력
frame_slide_num = tk.Frame(root)
frame_slide_num.pack(fill="x", padx=20, pady=(0, 10))

tk.Label(frame_slide_num, text="서식을 복사할 기준 슬라이드 번호:", font=NORMAL_FONT).pack(side="left")
# 1부터 100까지 선택 가능한 스핀박스 (기본값 1)
spin_slide = tk.Spinbox(frame_slide_num, from_=1, to=100, width=5, font=NORMAL_FONT)
spin_slide.pack(side="left", padx=5)

# 가사가 시작될 슬라이드 위치 입력
frame_insert_num = tk.Frame(root)
frame_insert_num.pack(fill="x", padx=20, pady=(0, 10))

tk.Label(frame_insert_num, text="가사가 시작될 삽입 위치 (슬라이드 번호):", font=NORMAL_FONT).pack(side="left")
# 1부터 100까지 선택 가능, 기본적으로 2번째 장부터 들어가도록 초기값 설정 가능
spin_insert = tk.Spinbox(frame_insert_num, from_=1, to=100, width=5, font=NORMAL_FONT)
spin_insert.pack(side="left", padx=5)

# (2) 곡 입력 창 (제목 + 가사)
frame_input = tk.Frame(root)
frame_input.pack(fill="x", padx=20, pady=5)
tk.Label(frame_input, text="2. 노래 제목 및 가사 입력", font=TITLE_FONT).pack(anchor="w")

entry_title = tk.Entry(frame_input, font=NORMAL_FONT)
entry_title.pack(fill="x", pady=(5, 5))
entry_title.insert(0, "노래 제목을 입력하세요")
entry_title.bind("<FocusIn>", lambda args: entry_title.delete(0, 'end') if entry_title.get() == "노래 제목을 입력하세요" else None)

scrollbar_text = tk.Scrollbar(frame_input)
scrollbar_text.pack(side="right", fill="y")
text_lyrics = tk.Text(frame_input, yscrollcommand=scrollbar_text.set, font=NORMAL_FONT, height=8)
text_lyrics.pack(fill="x", pady=5)
scrollbar_text.config(command=text_lyrics.yview)

btn_add = tk.Button(frame_input, text="추가하기 ⬇", bg="#e0e0e0", font=TITLE_FONT, command=add_song)
btn_add.pack(fill="x", pady=5)

# (3) 추가된 노래 목록
frame_list = tk.Frame(root)
frame_list.pack(fill="both", expand=True, padx=20, pady=5)
tk.Label(frame_list, text="3. 생성할 노래 목록", font=TITLE_FONT).pack(anchor="w")

scrollbar_list = tk.Scrollbar(frame_list)
scrollbar_list.pack(side="right", fill="y")
listbox_songs = tk.Listbox(frame_list, yscrollcommand=scrollbar_list.set, font=NORMAL_FONT, height=6, selectmode=tk.EXTENDED)
listbox_songs.pack(fill="both", expand=True, pady=5)
scrollbar_list.config(command=listbox_songs.yview)

# 🎯 [수정된 부분] 버튼들을 가로로 나란히 배치하기 위한 하위 프레임 생성
frame_list_buttons = tk.Frame(frame_list)
frame_list_buttons.pack(fill="x", pady=2)

# 🎯 [왼쪽 영역] 삭제 버튼 및 순서 변경 버튼 배치
btn_remove = tk.Button(frame_list_buttons, text="선택 삭제", command=remove_song, font=NORMAL_FONT)
btn_remove.pack(side="left")

btn_up = tk.Button(frame_list_buttons, text="▲ 위로", command=move_song_up, font=NORMAL_FONT)
btn_up.pack(side="left", padx=(10, 2))

btn_down = tk.Button(frame_list_buttons, text="▼ 아래로", command=move_song_down, font=NORMAL_FONT)
btn_down.pack(side="left")

# 🎯 [오른쪽 영역] DB 관련 버튼들 (오른쪽부터 차례대로 쌓임)
btn_save = tk.Button(frame_list_buttons, text="새 파일로 저장", command=save_song_list, font=NORMAL_FONT)
btn_save.pack(side="right")

btn_append = tk.Button(frame_list_buttons, text="기존 DB에 추가", command=append_to_song_list, font=NORMAL_FONT)
btn_append.pack(side="right", padx=(5, 0))

btn_load = tk.Button(frame_list_buttons, text="목록 불러오기", command=load_song_list, font=NORMAL_FONT)
btn_load.pack(side="right", padx=(5, 0))

# (4) 실행 버튼
frame_bottom = tk.Frame(root)
frame_bottom.pack(fill="x", padx=20, pady=(5, 15))

btn_generate = tk.Button(frame_bottom, text="전체 곡 PPT 만들기 및 저장", bg="#4CAF50", fg="white", 
                         font=("맑은 고딕", 12, "bold"), height=2, command=on_click_generate)
btn_generate.pack(fill="x")

lbl_status = tk.Label(frame_bottom, text="준비됨", font=NORMAL_FONT, fg="gray")
lbl_status.pack(pady=(5, 0))

root.mainloop()